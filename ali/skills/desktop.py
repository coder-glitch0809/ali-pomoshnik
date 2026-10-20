"""Dastur ichida ishlash (oynalar, sichqoncha, klaviatura) va fayllarni o'qib tahlil qilish.

Ish tartibi: take_screenshot -> focus_window / mouse_click / type_text / press_keys -> take_screenshot (tekshirish).
mouse_click koordinatalari oxirgi ekran rasmidagi koordinatalar - haqiqiy ekranga avtomatik moslashtiriladi.
"""
import os
import platform
import subprocess
import time
from datetime import datetime
from pathlib import Path

from . import skill

_scale = (1.0, 1.0)
_dpi_done = False
MAX_TEXT = 30000


def dpi_aware():
    """Windows'da ekran rasmi va sichqoncha bir xil (fizik) piksellarda ishlashi uchun."""
    global _dpi_done
    if _dpi_done or platform.system() != "Windows":
        return
    _dpi_done = True
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except Exception:
        pass


def set_scale(full: tuple[int, int], shown: tuple[int, int]):
    global _scale
    _scale = (full[0] / max(shown[0], 1), full[1] / max(shown[1], 1))


def _gui():
    dpi_aware()
    try:
        import pyautogui
    except ImportError as e:
        raise RuntimeError("pyautogui o'rnatilmagan: pip install pyautogui") from e
    pyautogui.FAILSAFE = True  # sichqonchani ekran burchagiga olib borsangiz - to'xtaydi
    pyautogui.PAUSE = 0.15
    return pyautogui


def _windows():
    import pygetwindow
    return [w for w in pygetwindow.getAllWindows() if w.title.strip() and w.width > 50 and w.height > 50]


@skill(description="Ochiq oynalar ro'yxati (sarlavha, joylashuv, o'lcham, faol/yo'q). Dastur ichida ishlashdan oldin ishlat.")
def list_windows() -> str:
    rows = [f"- {w.title} | x={w.left} y={w.top} {w.width}x{w.height}{' | FAOL' if w.isActive else ''}"
            f"{' | yig`ilgan' if w.isMinimized else ''}" for w in _windows()]
    return "\n".join(rows[:80]) or "Ochiq oyna topilmadi"


@skill(
    description="Oynani old planga chiqaradi va faollashtiradi (sarlavhaning bir qismi bo'yicha qidiradi).",
    params={"title": {"type": "string", "description": "Oyna sarlavhasi yoki uning bir qismi, masalan 'Excel'"}},
    dangerous=True,
)
def focus_window(title: str) -> str:
    matches = [w for w in _windows() if title.lower() in w.title.lower()]
    if not matches:
        return f"'{title}' oynasi topilmadi. list_windows bilan ro'yxatni ko'r."
    w = matches[0]
    if w.isMinimized:
        w.restore()
    try:
        w.activate()
    except Exception:
        _gui().click(w.left + w.width // 2, w.top + 10)  # activate ishlamasa - sarlavha satriga bosish
    time.sleep(0.4)
    return f"Faollashtirildi: {w.title}"


@skill(
    description=("Sichqoncha bilan bosadi. x, y - OXIRGI take_screenshot rasmidagi koordinatalar (avtomatik moslashtiriladi). "
                 "Bosishdan keyin natijani take_screenshot bilan tekshir."),
    params={
        "x": {"type": "integer", "description": "Rasmdagi X"},
        "y": {"type": "integer", "description": "Rasmdagi Y"},
        "button": {"type": "string", "enum": ["left", "right", "middle"], "description": "Tugma (standart: left)"},
        "double": {"type": "boolean", "description": "Ikki marta bosish"},
    },
    required=["x", "y"],
    dangerous=True,
)
def mouse_click(x: int, y: int, button: str = "left", double: bool = False) -> str:
    gui = _gui()
    rx, ry = round(x * _scale[0]), round(y * _scale[1])
    gui.click(rx, ry, clicks=2 if double else 1, button=button)
    return f"Bosildi: {button}{' x2' if double else ''} ({rx}, {ry})"


@skill(
    description="Sichqoncha g'ildiragi bilan aylantiradi (musbat - yuqoriga, manfiy - pastga).",
    params={"amount": {"type": "integer", "description": "Masalan -5 (pastga) yoki 5 (yuqoriga)"}},
    dangerous=True,
)
def mouse_scroll(amount: int) -> str:
    _gui().scroll(amount * 120)
    return f"Aylantirildi: {amount}"


def _set_clipboard(text: str):
    if platform.system() == "Windows":
        subprocess.run(["powershell", "-NoProfile", "-Command",
                        "[Console]::InputEncoding=[Text.Encoding]::UTF8; [Console]::In.ReadToEnd() | Set-Clipboard"],
                       input=text, text=True, encoding="utf-8", timeout=20, check=True)
    else:
        raise RuntimeError("Lotin bo'lmagan matnni yozish faqat Windows'da")


@skill(
    description=("Faol oynaga matn yozadi (klaviaturadan terilgandek). Kirill/o'zbek harflari ham ishlaydi. "
                 "Avval kerakli maydonni mouse_click bilan tanla. enter=true bo'lsa oxirida Enter bosiladi."),
    params={"text": {"type": "string", "description": "Yoziladigan matn"},
            "enter": {"type": "boolean", "description": "Oxirida Enter bosish"}},
    required=["text"],
    dangerous=True,
)
def type_text(text: str, enter: bool = False) -> str:
    gui = _gui()
    if text.isascii() and "\n" not in text:
        gui.write(text, interval=0.01)
    else:  # kirill va ko'p qatorli matn - bufer orqali
        _set_clipboard(text)
        gui.hotkey("ctrl", "v")
    if enter:
        gui.press("enter")
    return f"Yozildi: {len(text)} belgi" + (" + Enter" if enter else "")


@skill(
    description=("Klaviatura tugmasi yoki kombinatsiyasi. Misollar: 'enter', 'esc', 'tab', 'ctrl+s', 'ctrl+shift+esc', "
                 "'alt+f4', 'win+r', 'ctrl+a'. times - necha marta."),
    params={"keys": {"type": "string", "description": "Tugmalar, '+' bilan"},
            "times": {"type": "integer", "description": "Necha marta (standart 1)"}},
    required=["keys"],
    dangerous=True,
)
def press_keys(keys: str, times: int = 1) -> str:
    gui = _gui()
    parts = [k.strip().lower().replace("win", "winleft") for k in keys.split("+") if k.strip()]
    for _ in range(max(1, min(times, 50))):
        gui.hotkey(*parts) if len(parts) > 1 else gui.press(parts[0])
    return f"Bosildi: {keys}" + (f" x{times}" if times > 1 else "")


@skill(
    description="Papkadagi fayl va papkalar ro'yxati (o'lcham, o'zgargan sana). pattern bilan filtrlash, recursive - ichki papkalar ham.",
    params={"path": {"type": "string", "description": "Papka yo'li, masalan C:/Users/Abbos/Desktop yoki ~/Documents"},
            "pattern": {"type": "string", "description": "Masalan *.docx (standart: *)"},
            "recursive": {"type": "boolean", "description": "Ichki papkalarni ham ko'rish"}},
    required=["path"],
)
def list_folder(path: str, pattern: str = "*", recursive: bool = False) -> str:
    root = Path(os.path.expandvars(path)).expanduser()
    if not root.is_dir():
        return f"Papka topilmadi: {root}"
    items = sorted(root.rglob(pattern) if recursive else root.glob(pattern), key=lambda p: (p.is_file(), p.name.lower()))
    rows = []
    for p in items[:400]:
        try:
            st = p.stat()
        except OSError:
            continue
        when = datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M")
        rel = p.relative_to(root)
        rows.append(f"[papka] {rel}" if p.is_dir() else f"{rel}  ({st.st_size / 1024:.1f} KB, {when})")
    more = f"\n... yana {len(items) - 400} ta" if len(items) > 400 else ""
    return f"{root}:\n" + ("\n".join(rows) or "(bo'sh)") + more


def _read_docx(p):
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for i, table in enumerate(d.tables, 1):
        parts.append(f"\n[Jadval {i}]")
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
    return "\n".join(parts)


def _read_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n[Varaq: {ws.title}]")
        for row in ws.iter_rows(max_row=500, values_only=True):
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _read_pptx(p):
    from pptx import Presentation
    parts = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        parts.append(f"\n[Slayd {i}]\n" + "\n".join(texts))
    return "\n".join(parts)


def _read_pdf(p):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError("PDF uchun: pip install pypdf") from e
    return "\n".join(f"\n[Sahifa {i}]\n{page.extract_text() or ''}" for i, page in enumerate(PdfReader(str(p)).pages, 1))


def _read_text(p):
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


READERS = {".docx": _read_docx, ".xlsx": _read_xlsx, ".xlsm": _read_xlsx, ".pptx": _read_pptx, ".pdf": _read_pdf}
IMAGES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}


@skill(
    description=("Istalgan faylni o'qiydi va tahlil uchun matnini qaytaradi: Word (.docx), Excel (.xlsx), PowerPoint (.pptx), "
                 "PDF, rasm, kod, txt, csv, json, log va h.k. Katta fayl - offset bilan davomini o'qi."),
    params={"path": {"type": "string", "description": "Fayl yo'li"},
            "offset": {"type": "integer", "description": "Shu belgidan boshlab o'qish (davomi uchun)"}},
    required=["path"],
)
def read_file(path: str, offset: int = 0):
    p = Path(os.path.expandvars(path)).expanduser()
    if not p.is_file():
        return f"Fayl topilmadi: {p}"
    ext = p.suffix.lower()
    if ext in IMAGES:
        import base64
        from PIL import Image
        import io
        img = Image.open(p)
        img.thumbnail((1568, 1568))
        buf = io.BytesIO()
        img.convert("RGB").save(buf, format="JPEG", quality=85)
        return [{"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                             "data": base64.b64encode(buf.getvalue()).decode()}},
                {"type": "text", "text": f"Rasm: {p.name} ({img.width}x{img.height})"}]
    text = READERS.get(ext, _read_text)(p)
    chunk = text[offset:offset + MAX_TEXT]
    tail = (f"\n\n... davomi bor: read_file offset={offset + MAX_TEXT} (jami {len(text)} belgi)"
            if len(text) > offset + MAX_TEXT else "")
    return f"{p} ({len(text)} belgi):\n{chunk}{tail}"


@skill(
    description="Fayl yoki papkani o'z dasturida ochadi (docx - Word'da, papka - Explorer'da, .exe - ishga tushadi).",
    params={"path": {"type": "string", "description": "Fayl yoki papka yo'li"}},
    dangerous=True,
)
def open_path(path: str) -> str:
    p = Path(os.path.expandvars(path)).expanduser()
    if not p.exists():
        return f"Topilmadi: {p}"
    if platform.system() == "Windows":
        os.startfile(str(p))
    elif platform.system() == "Darwin":
        subprocess.Popen(["open", str(p)])
    else:
        subprocess.Popen(["xdg-open", str(p)])
    return f"Ochildi: {p}"
