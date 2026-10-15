"""Office fayllarini serverning o'zida yaratish: PowerPoint (.pptx), Word (.docx), Excel (.xlsx).

Agent faqat mazmunni tuzilgan ko'rinishda beradi, dizayn shu yerda: 16:9, yagona rang mavzusi, sarlavha va bo'lim
slaydlari, matn hajmiga qarab shrift, sig'masa davom slaydi, jadval, diagramma, raqamli kartalar, izohlar.
Fayl outbox'ga yoziladi va so'rov tugagach Telegram/saytga avtomatik yuboriladi. Skript yozish yoki tasdiq shart emas.
"""
from __future__ import annotations

import re

from .. import outbox
from . import skill

THEMES = {
    # primary, dark, accent, light background, text
    "blue": ("1F4E99", "0B2545", "F2A541", "F4F7FC", "1E2A3A"),
    "green": ("1B7F5B", "0C3B2E", "F4C430", "F3FAF6", "1E2E28"),
    "purple": ("5B2A86", "2A1442", "F28F3B", "F7F3FB", "2A2233"),
    "red": ("B23A48", "4A1220", "F6BD60", "FCF5F5", "2E1E22"),
    "orange": ("D9631E", "4A2208", "2A9D8F", "FDF6F0", "2E241E"),
    "dark": ("3A86FF", "0D1B2A", "FFBE0B", "1B263B", "E0E6ED"),
}
SLIDE_TYPES = ["bullets", "two_column", "table", "chart", "stats", "section", "quote"]
MAX_BULLETS = 7


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_color)


def _clean(value, limit: int = 600) -> str:
    return " ".join(str(value or "").split())[:limit]


def _list(value, limit: int = 40) -> list[str]:
    if isinstance(value, str):
        value = [v for v in re.split(r"\n+", value) if v.strip()]
    return [_clean(v, 400) for v in (value or []) if _clean(v)][:limit]


# ---------------------------------------------------------------- PowerPoint
class _Deck:
    def __init__(self, theme: str, title: str):
        from pptx import Presentation
        from pptx.util import Inches
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.primary, self.dark, self.accent, self.light, self.text = THEMES.get(theme, THEMES["blue"])
        self.dark_theme = theme == "dark"
        self.title = title
        self.blank = self.prs.slide_layouts[6]

    # -- primitives
    def _slide(self, background: str | None = None):
        slide = self.prs.slides.add_slide(self.blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(background or self.light)
        return slide

    def _box(self, slide, x, y, w, h, color: str, shape=None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        s = slide.shapes.add_shape(shape or MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(color)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def _text(self, slide, x, y, w, h, text: str, size: int, color: str, bold=False, align="left", anchor="top"):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for i, line in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            run = p.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
            run.font.color.rgb = _rgb(color)
        return tb

    def _bullets(self, slide, x, y, w, h, items: list[str], size: int, color: str):
        from pptx.util import Inches, Pt
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(size * (1.1 if len(items) <= 4 else 0.6))
            dot = p.add_run()
            dot.text = "●  "
            dot.font.size, dot.font.color.rgb, dot.font.name = Pt(size * 0.7), _rgb(self.accent), "Calibri"
            head, sep, rest = item.partition(": ")
            if sep and len(head) <= 40:  # "Термин: пояснение" -> термин жирным
                run = p.add_run()
                run.text = head + ": "
                run.font.bold = True
                run.font.size, run.font.color.rgb, run.font.name = Pt(size), _rgb(color), "Calibri"
                item = rest
            run = p.add_run()
            run.text = item
            run.font.size, run.font.color.rgb, run.font.name = Pt(size), _rgb(color), "Calibri"
        return tb

    def _font_for(self, items: list[str], base: int = 24) -> int:
        chars = sum(len(i) for i in items)
        if chars < 160 and len(items) <= 4:
            return base + 4
        return base if chars < 280 else base - 4 if chars < 480 else base - 6 if chars < 700 else base - 8

    def _header(self, slide, title: str, number: int):
        fg = self.text
        self._box(slide, 0, 0, 13.333, 0.12, self.primary)
        self._box(slide, 0.6, 0.55, 0.12, 0.75, self.accent)
        self._text(slide, 0.9, 0.42, 11.6, 1.0, title, 32 if len(title) < 45 else 26, fg, bold=True, anchor="middle")
        muted = "8A94A6"
        self._text(slide, 0.6, 6.95, 9, 0.4, self.title, 11, muted)
        self._text(slide, 11.7, 6.95, 1.0, 0.4, str(number), 11, muted, align="right")

    def _notes(self, slide, notes: str):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    # -- slide types
    def title_slide(self, title: str, subtitle: str, author: str):
        from pptx.enum.shapes import MSO_SHAPE
        slide = self._slide(self.dark)
        self._box(slide, 8.9, -1.6, 6.5, 6.5, self.primary, MSO_SHAPE.OVAL)
        self._box(slide, 10.6, 4.4, 4.0, 4.0, self.accent, MSO_SHAPE.OVAL)
        self._box(slide, 0.8, 2.15, 0.15, 2.1, self.accent)
        self._text(slide, 1.2, 1.9, 9.0, 2.0, title, 46 if len(title) < 40 else 36, "FFFFFF", bold=True, anchor="bottom")
        if subtitle:
            self._text(slide, 1.2, 4.0, 8.5, 1.2, subtitle, 22, "D6DEEB")
        if author:
            self._text(slide, 1.2, 6.3, 8, 0.5, author, 14, "AFC0D6")

    def section(self, title: str, subtitle: str, number: int):
        slide = self._slide(self.primary)
        self._text(slide, 1.0, 2.0, 2.5, 2.0, f"{number:02d}", 80, self.accent, bold=True)
        self._text(slide, 3.4, 2.3, 9.0, 1.4, title, 40, "FFFFFF", bold=True, anchor="middle")
        if subtitle:
            self._text(slide, 3.4, 3.8, 9.0, 1.2, subtitle, 20, "E6ECF5")
        return slide

    def bullets(self, title: str, items: list[str], number: int):
        slide = self._slide()
        self._header(slide, title, number)
        size = self._font_for(items)
        self._box(slide, 0.6, 1.65, 12.1, 5.05, "FFFFFF" if not self.dark_theme else self.dark)
        self._bullets(slide, 0.95, 1.9, 11.4, 4.7, items, size, self.text)
        return slide

    def two_column(self, title: str, left: dict, right: dict, number: int):
        slide = self._slide()
        self._header(slide, title, number)
        for i, col in enumerate((left, right)):
            x = 0.6 + i * 6.15
            card = "FFFFFF" if not self.dark_theme else self.dark
            self._box(slide, x, 1.65, 5.95, 5.05, card)
            self._box(slide, x, 1.65, 5.95, 0.75, self.primary if i == 0 else self.accent)
            self._text(slide, x + 0.3, 1.7, 5.4, 0.65, col.get("title", ""), 20, "FFFFFF", bold=True, anchor="middle")
            items = _list(col.get("bullets"), MAX_BULLETS)
            self._bullets(slide, x + 0.3, 2.6, 5.4, 4.0, items, self._font_for(items, 22), self.text)
        return slide

    def table(self, title: str, rows: list[list], number: int):
        from pptx.util import Inches, Pt
        slide = self._slide()
        self._header(slide, title, number)
        rows = [[_clean(c, 120) for c in r[:8]] for r in rows[:13]]
        cols = max(len(r) for r in rows)
        rows = [r + [""] * (cols - len(r)) for r in rows]
        height = min(5.0, (0.7 if len(rows) <= 5 else 0.5) * len(rows))
        shape = slide.shapes.add_table(len(rows), cols, Inches(0.6), Inches(1.75), Inches(12.1), Inches(height))
        size = 22 if len(rows) <= 5 and cols <= 3 else 18 if len(rows) <= 7 and cols <= 4 else 14 if len(rows) <= 10 else 12
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = shape.table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(self.primary if r == 0 else ("FFFFFF" if r % 2 else "EEF2F8")
                                              if not self.dark_theme else (self.primary if r == 0 else self.dark))
                cell.text = value
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size, run.font.bold = Pt(size), r == 0
                        run.font.color.rgb = _rgb("FFFFFF" if r == 0 or self.dark_theme else self.text)
        return slide

    def chart(self, title: str, spec: dict, number: int):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt
        slide = self._slide()
        self._header(slide, title, number)
        kinds = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}
        kind = spec.get("type", "column") if spec.get("type") in kinds else "column"
        data = CategoryChartData()
        data.categories = [_clean(c, 40) for c in spec.get("categories", [])][:20]
        series = spec.get("series") or []
        if kind in ("pie", "doughnut"):
            series = series[:1]
        for s in series[:6]:
            values = [float(v) if isinstance(v, (int, float)) else 0.0 for v in s.get("values", [])][:len(data.categories)]
            data.add_series(_clean(s.get("name", ""), 40) or "Qator", values)
        width = 7.6 if spec.get("caption") else 12.1
        frame = slide.shapes.add_chart(kinds[kind], Inches(0.6), Inches(1.7), Inches(width), Inches(5.0), data)
        chart = frame.chart
        chart.font.size, chart.font.name = Pt(13), "Calibri"
        chart.font.color.rgb = _rgb(self.text)
        chart.has_title = False  # slayd sarlavhasi yetarli
        chart.has_legend = len(series) > 1 or kind in ("pie", "doughnut")
        if chart.has_legend:
            chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        palette = [self.primary, self.accent, "2A9D8F", "E76F51", "8D99AE", "6A4C93"]
        if kind in ("pie", "doughnut"):
            for i, point in enumerate(chart.plots[0].series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
            chart.plots[0].has_data_labels = True
            chart.plots[0].data_labels.number_format, chart.plots[0].data_labels.number_format_is_linked = '0%', False
            chart.plots[0].data_labels.show_percentage, chart.plots[0].data_labels.show_value = True, False
        else:
            for i, s in enumerate(chart.plots[0].series):
                fmt = s.format.line if kind == "line" else s.format.fill
                if kind == "line":
                    fmt.color.rgb, fmt.width = _rgb(palette[i % len(palette)]), Pt(3)
                else:
                    fmt.solid()
                    fmt.fore_color.rgb = _rgb(palette[i % len(palette)])
        if spec.get("caption"):
            self._box(slide, 8.5, 1.7, 4.2, 5.0, "FFFFFF" if not self.dark_theme else self.dark)
            self._text(slide, 8.8, 1.95, 3.7, 4.5, _clean(spec["caption"], 500), 17, self.text)
        return slide

    def stats(self, title: str, items: list[dict], number: int):
        slide = self._slide()
        self._header(slide, title, number)
        items = items[:4]
        width = (12.1 - 0.3 * (len(items) - 1)) / max(1, len(items))
        for i, item in enumerate(items):
            x = 0.6 + i * (width + 0.3)
            self._box(slide, x, 2.0, width, 4.0, "FFFFFF" if not self.dark_theme else self.dark)
            self._box(slide, x, 2.0, width, 0.12, self.accent if i % 2 else self.primary)
            self._text(slide, x + 0.2, 2.5, width - 0.4, 1.6, _clean(item.get("value"), 14), 54, self.primary,
                       bold=True, align="center", anchor="middle")
            self._text(slide, x + 0.2, 4.2, width - 0.4, 1.6, _clean(item.get("label"), 120), 18, self.text, align="center")
        return slide

    def quote(self, text: str, author: str):
        slide = self._slide(self.dark)
        self._text(slide, 0.9, 0.6, 2, 2, "“", 160, self.accent, bold=True)
        self._text(slide, 1.6, 2.0, 10.2, 3.2, text, 34 if len(text) < 140 else 26, "FFFFFF", anchor="middle")
        if author:
            self._text(slide, 1.6, 5.4, 10.2, 0.6, "— " + author, 18, self.accent)
        return slide

    def ending(self, text: str):
        from pptx.enum.shapes import MSO_SHAPE
        slide = self._slide(self.dark)
        self._box(slide, -1.5, 4.0, 5.0, 5.0, self.primary, MSO_SHAPE.OVAL)
        self._text(slide, 1.0, 2.4, 11.3, 1.6, text, 48, "FFFFFF", bold=True, align="center", anchor="middle")
        self._text(slide, 1.0, 4.0, 11.3, 0.6, self.title, 18, "AFC0D6", align="center")


@skill(
    description=(
        "Chiroyli, tahrirlanadigan PowerPoint (.pptx) taqdimot yaratadi (16:9, yagona rang mavzusi, sarlavha slaydi, "
        "bo'limlar, punktlar, ikki ustun, jadval, diagramma, raqamli kartalar, iqtibos, yakuniy slayd, izohlar). "
        "Fayl so'rov tugagach foydalanuvchiga AVTOMATIK yuboriladi (Telegram'da hujjat). Mazmunni o'zing to'liq yoz: "
        "har slaydda 3-6 ta aniq punkt, 8-14 slayd. Slayd turlari: bullets {title, bullets}, two_column {title, left: "
        "{title, bullets}, right: {title, bullets}}, table {title, rows: [[sarlavhalar], [qator]...]}, chart {title, chart: "
        "{type: column|bar|line|pie|doughnut, categories, series: [{name, values}], caption}}, stats {title, stats: "
        "[{value, label}]}, section {title, subtitle}, quote {text, author}. Har slaydga notes (ma'ruzachi izohi) qo'sh."
    ),
    params={
        "title": {"type": "string", "description": "Taqdimot nomi"},
        "subtitle": {"type": "string", "description": "Sarlavha slaydidagi kichik matn"},
        "author": {"type": "string", "description": "Muallif / maktab / sana (ixtiyoriy)"},
        "theme": {"type": "string", "enum": list(THEMES), "description": "Rang mavzusi (standart blue)"},
        "slides": {"type": "array", "items": {"type": "object"}, "description": "Slaydlar ro'yxati (type + maydonlar)"},
        "ending": {"type": "string", "description": "Yakuniy slayd matni, masalan 'Спасибо за внимание!' (bo'sh = yo'q)"},
    },
    required=["title", "slides"],
)
def make_presentation(title: str, slides: list, subtitle: str = "", author: str = "", theme: str = "blue",
                      ending: str = "") -> str:
    from pptx import Presentation
    title = _clean(title, 120) or "Taqdimot"
    if not isinstance(slides, list) or not slides:
        return "slides bo'sh: kamida bitta slayd kerak."
    deck = _Deck(theme if theme in THEMES else "blue", title)
    deck.title_slide(title, _clean(subtitle, 200), _clean(author, 120))
    sections = 0
    for raw in slides[:40]:
        if not isinstance(raw, dict):
            continue
        kind = raw.get("type") if raw.get("type") in SLIDE_TYPES else ("table" if raw.get("rows") else "chart"
                if raw.get("chart") else "stats" if raw.get("stats") else "bullets")
        head = _clean(raw.get("title"), 120)
        number = len(deck.prs.slides) + 1
        if kind == "section":
            sections += 1
            slide = deck.section(head, _clean(raw.get("subtitle"), 200), sections)
        elif kind == "two_column":
            slide = deck.two_column(head, raw.get("left") or {}, raw.get("right") or {}, number)
        elif kind == "table" and isinstance(raw.get("rows"), list) and raw["rows"] and all(isinstance(r, list) for r in raw["rows"]):
            slide = deck.table(head, raw["rows"], number)
        elif kind == "chart" and isinstance(raw.get("chart"), dict) and raw["chart"].get("categories"):
            slide = deck.chart(head, raw["chart"], number)
        elif kind == "stats" and isinstance(raw.get("stats"), list) and raw["stats"]:
            slide = deck.stats(head, [s for s in raw["stats"] if isinstance(s, dict)], number)
        elif kind == "quote":
            slide = deck.quote(_clean(raw.get("text") or head, 400), _clean(raw.get("author"), 80))
        else:
            items = _list(raw.get("bullets") or raw.get("text"))
            if not items:
                continue
            # Ko'p punkt sig'masa - davom slaydi
            for part in range(0, len(items), MAX_BULLETS):
                slide = deck.bullets(head + (" (davomi)" if part else ""), items[part:part + MAX_BULLETS],
                                     len(deck.prs.slides) + 1)
                deck._notes(slide, _clean(raw.get("notes"), 2000))
            continue
        deck._notes(slide, _clean(raw.get("notes"), 2000))
    if ending:
        deck.ending(_clean(ending, 80))
    path = outbox.new_path(title, "pptx")
    deck.prs.save(path)
    count = len(Presentation(path).slides)  # qayta ochib tekshirish
    outbox.register(path)
    return (f"✅ PowerPoint tayyor: {path.name} ({count} slayd, mavzu {theme}). Fayl javob bilan birga avtomatik "
            "yuboriladi - havola to'qima, faqat nima qilinganini qisqa ayt.")


# ---------------------------------------------------------------- Word
@skill(
    description=(
        "Professional Word (.docx) hujjat yaratadi: sarlavha sahifasi, mundarija (ixtiyoriy), bo'lim sarlavhalari, "
        "paragraflar, punkt va raqamli ro'yxatlar, jadvallar, sahifa raqamlari. Referat, doklad, dars ishlanmasi, "
        "ariza, xat, hisobot, test uchun. Fayl foydalanuvchiga AVTOMATIK yuboriladi. sections: [{heading, level (1|2), "
        "paragraphs: [...], bullets: [...], numbered: [...], table: [[sarlavhalar], [qator]...]}]."
    ),
    params={
        "title": {"type": "string", "description": "Hujjat nomi"},
        "subtitle": {"type": "string", "description": "Kichik sarlavha / muallif / sana"},
        "sections": {"type": "array", "items": {"type": "object"}, "description": "Bo'limlar"},
        "toc": {"type": "boolean", "description": "Mundarija qo'shish (Word'da F9 bilan yangilanadi)"},
        "title_page": {"type": "boolean", "description": "Alohida sarlavha sahifasi (standart false)"},
    },
    required=["title", "sections"],
)
def make_word(title: str, sections: list, subtitle: str = "", toc: bool = False, title_page: bool = False) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    title = _clean(title, 150) or "Hujjat"
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name, styles["Normal"].font.size = "Calibri", Pt(12)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Title", 26)):
        styles[name].font.name, styles[name].font.size = "Calibri", Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string("1F4E99")

    def field(paragraph, code: str):
        run = paragraph.add_run()
        for tag, text in (("begin", None), (None, code), ("separate", None), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    head = doc.add_paragraph(title, style="Title")
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER if title_page else WD_ALIGN_PARAGRAPH.LEFT
    if subtitle:
        p = doc.add_paragraph(_clean(subtitle, 300))
        p.alignment = head.alignment
        p.runs[0].italic = True
    if title_page:
        doc.add_page_break()
    if toc:
        doc.add_paragraph("Mundarija / Содержание", style="Heading 1")
        field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
        doc.add_page_break()
    tables = 0
    for raw in (sections or [])[:60]:
        if not isinstance(raw, dict):
            continue
        if raw.get("heading"):
            doc.add_heading(_clean(raw["heading"], 200), level=2 if raw.get("level") == 2 else 1)
        for text in _list(raw.get("paragraphs"), 60):
            p = doc.add_paragraph(text)
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for text in _list(raw.get("bullets"), 60):
            doc.add_paragraph(text, style="List Bullet")
        for text in _list(raw.get("numbered"), 60):
            doc.add_paragraph(text, style="List Number")
        rows = raw.get("table")
        if isinstance(rows, list) and rows and all(isinstance(r, list) for r in rows):
            rows = [[_clean(c, 300) for c in r[:10]] for r in rows[:200]]
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c in range(cols):
                    table.cell(r, c).text = row[c] if c < len(row) else ""
            tables += 1
            doc.add_paragraph()
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, "PAGE")
    path = outbox.new_path(title, "docx")
    doc.save(path)
    paragraphs = len(Document(path).paragraphs)
    outbox.register(path)
    return (f"✅ Word tayyor: {path.name} ({paragraphs} paragraf, {tables} jadval). Fayl javob bilan birga avtomatik "
            "yuboriladi." + (" Mundarija Word'da ochilganda F9 bilan yangilanadi." if toc else ""))


# ---------------------------------------------------------------- Excel
_DANGEROUS_FORMULA = re.compile(r"\|| DDE|WEBSERVICE|FILTERXML|CALL\(|REGISTER\(|EXEC\(|HYPERLINK\(", re.I)


def _cell_value(value):
    if isinstance(value, str):
        value = value.strip()
        if value.startswith("="):
            return value if not _DANGEROUS_FORMULA.search(value) else "'" + value
        number = value.replace(" ", "").replace(",", ".")
        if re.fullmatch(r"-?\d+(\.\d+)?", number) and not (len(number) > 1 and number.startswith("0") and "." not in number):
            return float(number) if "." in number else int(number)
    return value


@skill(
    description=(
        "Chiroyli Excel (.xlsx) fayl yaratadi: bir nechta varaq, rangli sarlavha, filtr, muzlatilgan sarlavha, ustun kengligi, "
        "raqam formatlari, formulalar (\"=SUM(B2:B10)\" kabi), avtomatik JAMI qatori va diagramma. Jurnal, baholar, "
        "byudjet, hisobot, jadval uchun. Fayl foydalanuvchiga AVTOMATIK yuboriladi. sheets: [{name, columns: [...], "
        "rows: [[...]], totals: true|false, number_format: \"#,##0.00\", chart: {type: column|bar|line|pie, title, "
        "x_col: 1, y_cols: [2,3]}}]. Formulalar Excel ochilganda hisoblanadi."
    ),
    params={
        "title": {"type": "string", "description": "Fayl nomi"},
        "sheets": {"type": "array", "items": {"type": "object"}, "description": "Varaqlar"},
    },
    required=["title", "sheets"],
)
def make_excel(title: str, sheets: list) -> str:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    title = _clean(title, 120) or "Jadval"
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    used: set[str] = set()
    summary = []
    for index, raw in enumerate((sheets or [])[:20]):
        if not isinstance(raw, dict):
            continue
        name = re.sub(r"[\[\]:*?/\\]", "", _clean(raw.get("name"), 31)) or f"Varaq{index + 1}"
        while name in used:
            name = (name[:28] + f"_{index}")[:31]
        used.add(name)
        ws = wb.create_sheet(name)
        columns = [_clean(c, 80) for c in raw.get("columns") or []][:50]
        rows = [r for r in (raw.get("rows") or []) if isinstance(r, list)][:5000]
        width = max([len(columns)] + [len(r) for r in rows] + [1])
        if columns:
            ws.append(columns + [""] * (width - len(columns)))
        for r in rows:
            ws.append([_cell_value(v) for v in r[:width]])
        first_data = 2 if columns else 1
        last = ws.max_row
        if raw.get("totals") and rows and columns:
            total_row = last + 1
            ws.cell(total_row, 1, "JAMI / ИТОГО").font = Font(bold=True)
            for c in range(2, width + 1):
                if any(isinstance(ws.cell(r, c).value, (int, float)) for r in range(first_data, last + 1)):
                    letter = get_column_letter(c)
                    ws.cell(total_row, c, f"=SUM({letter}{first_data}:{letter}{last})").font = Font(bold=True)
            for c in range(1, width + 1):
                ws.cell(total_row, c).fill = PatternFill("solid", fgColor="FFF2CC")
        fmt = _clean(raw.get("number_format"), 30)
        for row in ws.iter_rows(min_row=first_data, max_row=ws.max_row, max_col=width):
            for cell in row:
                cell.border = border
                if fmt and isinstance(cell.value, float):
                    cell.number_format = fmt
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
                elif isinstance(cell.value, int) and abs(cell.value) >= 1000:
                    cell.number_format = "#,##0"
        if columns:
            for c in range(1, width + 1):
                cell = ws.cell(1, c)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E99")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(width)}{last}"
            ws.row_dimensions[1].height = 28
        for c in range(1, width + 1):
            longest = max((len(str(ws.cell(r, c).value or "")) for r in range(1, min(ws.max_row, 300) + 1)), default=8)
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, longest + 3))
        chart_spec = raw.get("chart")
        if isinstance(chart_spec, dict) and columns and last >= first_data:
            kind = chart_spec.get("type", "column")
            chart = PieChart() if kind == "pie" else LineChart() if kind == "line" else BarChart()
            if isinstance(chart, BarChart):
                chart.type = "bar" if kind == "bar" else "col"
            chart.title = _clean(chart_spec.get("title"), 80) or name
            chart.height, chart.width = 9, 18
            x_col = int(chart_spec.get("x_col") or 1)
            y_cols = [int(c) for c in chart_spec.get("y_cols") or [2] if isinstance(c, (int, float)) and 1 <= c <= width]
            for c in (y_cols[:1] if kind == "pie" else y_cols[:6]):
                chart.add_data(Reference(ws, min_col=c, min_row=1, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=x_col, min_row=first_data, max_row=last))
            ws.add_chart(chart, f"{get_column_letter(width + 2)}2")
        summary.append(f"{name}: {len(rows)} qator")
    if not wb.sheetnames:
        return "sheets bo'sh: kamida bitta varaq kerak."
    path = outbox.new_path(title, "xlsx")
    wb.save(path)
    load_workbook(path)  # qayta ochib tekshirish
    outbox.register(path)
    return (f"✅ Excel tayyor: {path.name} ({'; '.join(summary)}). Formulalar Excel ochilganda hisoblanadi. "
            "Fayl javob bilan birga avtomatik yuboriladi.")
