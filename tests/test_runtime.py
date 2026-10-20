import asyncio
import importlib
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

state = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1] / "data")
os.environ.update(ALI_STATE_DIR=state.name, ANTHROPIC_API_KEY="", TELEGRAM_BOT_TOKEN="",
                  ACCESS_PASSWORD="test-password", AUTO_LEARN="false", WATCH_PROJECTS="false")

from fastapi.testclient import TestClient
from ali import agents, auth
from ali.brain import Brain
from ali.config import config
from ali.service import brain
from ali.telegram_bot import TelegramBot
from runtime import app
from gateway import app as gateway


def tearDownModule():
    import logging
    from ali.memory import memory
    memory.conn.close()
    logging.shutdown()
    state.cleanup()


class RuntimeTests(unittest.TestCase):
    def test_telegram_setup_rate_limit_does_not_block_and_is_cached(self):
        from types import SimpleNamespace
        from telegram.error import RetryAfter
        bot=TelegramBot();bot.public.set_meta('bot_setup',{})
        method=AsyncMock(side_effect=RetryAfter(3600))
        bot.app=SimpleNamespace(bot=SimpleNamespace(set_my_name=method))
        asyncio.run(bot.configure_bot('set_my_name',name='Ali'))
        asyncio.run(bot.configure_bot('set_my_name',name='Ali'))
        self.assertEqual(method.await_count,1)
        bot.public.set_meta('bot_setup',{})
        method.side_effect=None
        asyncio.run(bot.configure_bot('set_my_name',name='Ali'))
        asyncio.run(bot.configure_bot('set_my_name',name='Ali'))
        self.assertEqual(method.await_count,2)
        bot.public.set_meta('bot_setup',{})

    def test_miniapp_signed_owner_authentication(self):
        import hashlib,hmac,json,time
        from urllib.parse import urlencode
        def signed(uid=999,age=0):
            fields={'auth_date':str(int(time.time())-age),'user':json.dumps({'id':uid,'first_name':'Untrusted'}),'query_id':'test-query'}
            secret=hmac.new(b'WebAppData',b'test-token',hashlib.sha256).digest()
            fields['hash']=hmac.new(secret,'\n'.join(f'{k}={v}' for k,v in sorted(fields.items())).encode(),hashlib.sha256).hexdigest()
            return urlencode(fields)
        with patch.object(config,'telegram_token','test-token'),patch.object(config,'telegram_owner_id',999):
            self.assertEqual(self.client.get('/miniapp').status_code,200)
            for invalid in [signed(888),signed(age=600),signed(age=-90),signed()+'&user=%7B%7D',signed().replace('query_id=test-query','query_id=forged'),None,{'id':999}]:
                self.assertEqual(self.client.post('/api/telegram/login',json={'init_data':invalid}).status_code,403)
                self.assertEqual(self.client.get('/api/status').status_code,401)
            result=self.client.post('/api/telegram/login',json={'init_data':signed()})
            self.assertEqual(result.status_code,200);self.assertEqual(result.json()['name'],'Abbos')
            cookie=self.client.cookies.get('ali_token');self.assertTrue(auth.check_token(cookie))
            self.assertIn('HttpOnly',result.headers['set-cookie'])
            self.assertEqual(self.client.get('/api/status').status_code,200)
            with patch('ali.auth.time.time',return_value=time.time()+9*3600):self.assertFalse(auth.check_token(cookie))
            with patch.object(config,'telegram_owner_id',888):self.assertFalse(auth.check_token(cookie))

    def test_owner_menu_and_abbos_greeting(self):
        from types import SimpleNamespace
        bot=TelegramBot();bot.owner_id=999;bot.users={999}
        owner=[b.text for row in bot.keyboard(999).keyboard for b in row]
        guest=[b.text for row in bot.keyboard(111).keyboard for b in row]
        self.assertIn('⚙️ Sozlamalar',owner);self.assertIn('🖥 Mini App',owner)
        self.assertNotIn('⚙️ Sozlamalar',guest);self.assertNotIn('🖥 Mini App',guest)
        msg=SimpleNamespace(reply_text=AsyncMock())
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=SimpleNamespace(id=999,first_name='OtherName'),effective_message=msg)
        ctx=SimpleNamespace(bot=SimpleNamespace(set_chat_menu_button=AsyncMock()))
        asyncio.run(bot.public_start(update,ctx))
        self.assertIn('Abbos',msg.reply_text.call_args.args[0]);self.assertNotIn('OtherName',msg.reply_text.call_args.args[0])
        update.effective_user.id=111
        for text in ['⚙️ Sozlamalar','🖥 Mini App']:
            self.assertTrue(asyncio.run(bot.menu_text(update,ctx,text)))
            self.assertIn('faqat administrator',msg.reply_text.call_args.args[0])

    def test_connections_auth_validation_and_persistence(self):
        from ali import connections
        from types import SimpleNamespace
        self.assertEqual(self.client.post('/api/connections/github',json={'token':'not-a-real-token'}).status_code,401)
        self.login()
        self.assertEqual(self.client.post('/api/connections/github',headers={'Origin':'https://evil.example'},json={'token':'not-a-real-token'}).status_code,403)
        self.assertEqual(self.client.post('/api/connections/google',json={'token':'not-a-real-token'}).status_code,400)
        self.assertEqual(self.client.post('/api/connections/github',json={'token':123}).status_code,400)
        response=SimpleNamespace(status_code=401,json=lambda:{})
        with patch('ali.connections.httpx.AsyncClient.get',new=AsyncMock(return_value=response)),patch('ali.connections.set_key') as save:
            self.assertEqual(self.client.post('/api/connections/github',json={'token':'not-a-real-token'}).status_code,400)
            save.assert_not_called()
        response.status_code=200;response.json=lambda:{'login':'test'}
        with patch('ali.connections.httpx.AsyncClient.get',new=AsyncMock(return_value=response)),patch.object(connections,'ENV_FILE',Path(state.name)/'connection.env'),patch.object(config,'github_token',''),patch.dict(os.environ):
            self.assertEqual(self.client.post('/api/connections/github',json={'token':'not-a-real-token'}).status_code,200)
            self.assertEqual(config.github_token,'not-a-real-token')
            self.assertIn('GITHUB_TOKEN',connections.ENV_FILE.read_text())
            data=self.client.get('/api/permissions').json()
            self.assertNotIn('not-a-real-token',str(data))
            self.assertTrue(all(c['url'].startswith('https://') for c in data['connections']))

    def test_public_welcome_uses_native_menu_and_does_not_require_login(self):
        from types import SimpleNamespace
        bot=TelegramBot();bot.owner_id=999
        message=SimpleNamespace(reply_text=AsyncMock())
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=SimpleNamespace(id=108,first_name='Visitor'),effective_message=message)
        ctx=SimpleNamespace(bot=SimpleNamespace(set_chat_menu_button=AsyncMock()))
        asyncio.run(bot.public_start(update,ctx))
        self.assertTrue(bot.public.get(108)['subscribed'])
        self.assertIn(f'{config.free_daily_questions} ta',message.reply_text.call_args.args[0])
        self.assertNotIn('/login',message.reply_text.call_args.args[0])
        self.assertEqual(ctx.bot.set_chat_menu_button.call_args.kwargs['menu_button'].type,'commands')

    def test_public_trading_signal_shows_payment_without_ai_or_quota(self):
        from types import SimpleNamespace
        bot=TelegramBot();bot.owner_id=999
        message=SimpleNamespace(reply_text=AsyncMock())
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=SimpleNamespace(id=109,first_name='Trader'),effective_message=message,message=message)
        with patch.object(bot,'classify_guest',new=AsyncMock()) as classify,patch.object(bot,'guest_reply',new=AsyncMock()) as reply:
            asyncio.run(bot.public_process(update,SimpleNamespace(),'BTCUSDT signal bering'))
        classify.assert_not_called();reply.assert_not_called()
        self.assertIn(config.telegram_payment_text,message.reply_text.call_args.args[0])
        self.assertEqual(bot.public.get(109).get('count',0),0)

    def test_public_weather_only_once_per_day(self):
        from types import SimpleNamespace
        from datetime import datetime
        from ali.telegram_store import TZ
        bot=TelegramBot();uid=109
        data=bot.public.get(uid);data.update(weather=True,location={'lat':41,'lon':69},subscribed=False);bot.public.put(uid,data)
        bot.app=SimpleNamespace(bot=SimpleNamespace(send_message=AsyncMock()))
        class Clock:
            @staticmethod
            def now(tz):return datetime(2026,10,15,8,0,tzinfo=TZ)
        with patch('ali.telegram_public.datetime',Clock),patch.object(bot.public,'subscribers',side_effect=lambda:[(uid,bot.public.get(uid))]),patch.object(bot,'weather_text',new=AsyncMock(return_value='Weather')):
            asyncio.run(bot.public_tick());asyncio.run(bot.public_tick())
        self.assertEqual(bot.app.bot.send_message.await_count,1)

    def test_weather_callback_preserves_owner_menu(self):
        from types import SimpleNamespace
        bot=TelegramBot();bot.owner_id=999
        q=SimpleNamespace(from_user=SimpleNamespace(id=999),data='pub:weather:off',answer=AsyncMock(),message=SimpleNamespace(reply_text=AsyncMock()))
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=q.from_user,callback_query=q)
        asyncio.run(bot.public_button(update,None))
        keyboard=q.message.reply_text.call_args.kwargs['reply_markup']
        self.assertIn('⚙️ Sozlamalar',[b.text for row in keyboard.keyboard for b in row])

    def test_stop_removes_weather_location_and_notifications(self):
        from types import SimpleNamespace
        bot=TelegramBot();uid=110
        data=bot.public.get(uid);data.update(weather=True,subscribed=True,location={'lat':41,'lon':69});bot.public.put(uid,data)
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=SimpleNamespace(id=uid),effective_message=SimpleNamespace(reply_text=AsyncMock()))
        asyncio.run(bot.public_stop(update,None))
        saved=bot.public.get(uid)
        self.assertFalse(saved['weather']);self.assertFalse(saved['subscribed']);self.assertNotIn('location',saved)

    def test_public_telegram_routes_guests_without_owner_brain(self):
        from types import SimpleNamespace
        bot=TelegramBot();bot.owner_id=999;bot.users={999}
        update=SimpleNamespace(effective_chat=SimpleNamespace(id=101,type='private'),effective_user=SimpleNamespace(id=101,first_name='Guest'),message=SimpleNamespace(text='Salom'))
        with patch.object(bot,'menu_text',new=AsyncMock(return_value=False)),patch.object(bot,'public_process',new=AsyncMock()) as public,patch.object(bot,'process',new=AsyncMock()) as owner:
            asyncio.run(bot.on_text(update,None))
            public.assert_awaited_once();owner.assert_not_awaited()

    def test_public_telegram_location_does_not_replace_owner_location(self):
        from types import SimpleNamespace
        from ali.hub import hub
        bot=TelegramBot();bot.owner_id=999
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=SimpleNamespace(id=102),message=SimpleNamespace(location=SimpleNamespace(latitude=40,longitude=70,horizontal_accuracy=1)))
        with patch.object(hub,'last_location',{'lat':1,'lon':2}),patch.object(bot,'save_location',new=AsyncMock()):
            asyncio.run(bot.on_location(update,None))
            self.assertEqual(hub.last_location,{'lat':1,'lon':2})

    def test_public_requests_deduplicate_and_subscription_duration(self):
        from ali.telegram_store import PublicStore
        import time
        store=PublicStore()
        rid,fresh=store.request(103,'trading_week','test')
        self.assertTrue(fresh)
        self.assertEqual(store.request(103,'trading_week','again'),(rid,False))
        before=time.time();store.resolve(rid,True)
        expires=store.get(103)['grants']['trading']
        self.assertAlmostEqual(expires-before,7*86400,delta=2)
        self.assertIsNone(store.resolve(rid,True))
        self.assertEqual(store.get(104)['grants'],{})

    def test_guest_cannot_approve_own_request(self):
        from types import SimpleNamespace
        bot=TelegramBot();bot.owner_id=999
        rid,_=bot.public.request(105,'content','test')
        query=SimpleNamespace(from_user=SimpleNamespace(id=105),data=f'own:{rid}:yes',answer=AsyncMock())
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=query.from_user,callback_query=query)
        asyncio.run(bot.public_button(update,None))
        self.assertEqual(bot.public.get(105)['grants'],{})

    def test_public_quota_and_history_are_separate(self):
        from ali.telegram_store import PublicStore
        store=PublicStore()
        for _ in range(10):self.assertTrue(store.quota(106))
        self.assertFalse(store.quota(106))
        self.assertTrue(store.quota(107))
        data=store.get(106);data['history']=[{'role':'user','content':'private guest test'}];store.put(106,data)
        self.assertEqual(store.get(107)['history'],[])
        from ali.memory import memory
        self.assertEqual(memory.vault_list('private guest test')['total'],0)

    def test_free_and_pro_tiers_pick_model_and_image_limit(self):
        from types import SimpleNamespace
        from unittest.mock import MagicMock
        import time
        bot=TelegramBot();bot.owner_id=999
        rid,_=bot.public.request(120,'pro_month','pay');bot.public.resolve(rid,True)
        self.assertTrue(bot.public.is_pro(120));self.assertFalse(bot.public.is_pro(121))
        self.assertAlmostEqual(bot.public.get(120)['grants']['pro']-time.time(),30*86400,delta=5)
        calls=[]
        async def create(**kw):
            calls.append(kw);return SimpleNamespace(content=[SimpleNamespace(type='text',text='javob')])
        client=MagicMock();client.__aenter__=AsyncMock(return_value=SimpleNamespace(messages=SimpleNamespace(create=create)));client.__aexit__=AsyncMock(return_value=False)
        with patch('ali.telegram_public.anthropic.AsyncAnthropic',return_value=client):
            asyncio.run(bot.guest_reply(121,'salom','general',False))
            asyncio.run(bot.guest_reply(120,'batafsil tahlil','general',True))
        self.assertEqual(calls[0]['model'],config.free_model);self.assertEqual(calls[0]['output_config'],{'effort':'low'})
        self.assertEqual(calls[1]['model'],config.model);self.assertNotIn('output_config',calls[1])
        # Bepul: 3 ta rasm, 4-chisida Pro taklifi
        msg=SimpleNamespace(reply_text=AsyncMock(),reply_photo=AsyncMock())
        update=SimpleNamespace(effective_chat=SimpleNamespace(type='private'),effective_user=SimpleNamespace(id=121,first_name='G'),message=msg,effective_message=msg)
        ctx=SimpleNamespace(bot=SimpleNamespace(send_chat_action=AsyncMock()))
        with patch('ali.providers.can_generate_image',return_value=True),patch('ali.providers.generate_image',return_value=b'png') as gen:
            for _ in range(config.free_daily_images+1):asyncio.run(bot.guest_image(update,ctx,121,'mushuk'))
            self.assertEqual(gen.call_count,config.free_daily_images)
            self.assertIn('Pro',msg.reply_text.call_args.args[0])
            update.effective_user.id=120
            for _ in range(config.free_daily_images+1):asyncio.run(bot.guest_image(update,ctx,120,'mushuk'))
            self.assertEqual(gen.call_count,2*config.free_daily_images+1)

    def test_full_access_mode_auto_approves_but_keeps_owner_deny(self):
        from ali import permissions
        from unittest.mock import Mock
        with patch.object(permissions,'load',return_value={'rules':{'run_powershell':'deny'}}):
            with patch.object(config,'full_access',False):
                self.assertEqual(permissions.mode('mouse_click',True),'ask')
                self.assertEqual(permissions.mode('run_powershell',True),'deny')
            with patch.object(config,'full_access',True):
                self.assertEqual(permissions.mode('mouse_click',True),'auto')
                self.assertEqual(permissions.mode('open_application'),'auto')
                self.assertEqual(permissions.mode('run_powershell',True),'deny')
        with patch.object(permissions,'load',return_value={'rules':{},'unavailable':True}),patch.object(config,'full_access',True):
            self.assertEqual(permissions.mode('mouse_click',True),'deny')
        # Render'da FULL_ACCESS yoqilmaydi
        from ali.config import Config
        with patch.dict(os.environ,{'FULL_ACCESS':'true','RENDER':'true'}):self.assertFalse(Config().full_access)
        with patch.dict(os.environ,{'FULL_ACCESS':'true'}):
            os.environ.pop('RENDER',None);self.assertTrue(Config().full_access)
        confirm=Mock(return_value=False);b=Brain(confirm=confirm)
        with patch.object(config,'full_access',True),patch.object(permissions,'load',return_value={'rules':{}}),\
             patch.object(b.registry['list_windows'],'func',return_value='oyna'):
            self.assertEqual(b._run_one('list_windows',{}),('oyna',False))
        confirm.assert_not_called()

    def test_read_file_and_list_folder_handle_office_and_text(self):
        from ali.skills import desktop
        from unittest.mock import Mock
        import docx,openpyxl
        root=Path(state.name)/'files';root.mkdir(exist_ok=True)
        d=docx.Document();d.add_paragraph('Salom hujjat');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='A';t.cell(0,1).text='B';d.save(root/'a.docx')
        wb=openpyxl.Workbook();wb.active.append(['narx',42]);wb.save(root/'b.xlsx')
        (root/'c.txt').write_bytes('Привет мир'.encode('cp1251'))
        self.assertIn('Salom hujjat',desktop.read_file(str(root/'a.docx')));self.assertIn('A | B',desktop.read_file(str(root/'a.docx')))
        self.assertIn('narx | 42',desktop.read_file(str(root/'b.xlsx')))
        self.assertIn('Привет мир',desktop.read_file(str(root/'c.txt')))
        self.assertIn('topilmadi',desktop.read_file(str(root/'yoq.txt')))
        listing=desktop.list_folder(str(root),'*.docx');self.assertIn('a.docx',listing);self.assertNotIn('b.xlsx',listing)
        (root/'big.txt').write_text('x'*(desktop.MAX_TEXT+10))
        self.assertIn('offset=',desktop.read_file(str(root/'big.txt')))
        desktop.set_scale((3000,2000),(1500,1000))
        fake=Mock()
        with patch.object(desktop,'_gui',return_value=fake):
            desktop.mouse_click(100,50);fake.click.assert_called_with(200,100,clicks=1,button='left')
            desktop.press_keys('ctrl+s');fake.hotkey.assert_called_with('ctrl','s')

    def test_consult_experts_uses_connected_models(self):
        from ali import providers
        from ali.skills.experts import consult_experts
        with patch.object(config,'openai_api_key',''),patch.object(config,'gemini_api_key',''):
            self.assertIn('ulanmagan',consult_experts('savol'))
        with patch.object(config,'openai_api_key','k'),patch.object(config,'gemini_api_key',''),\
             patch.object(providers,'ask_openai',return_value='GPT fikri'):
            self.assertIn('GPT fikri',consult_experts('savol'))

    def test_malformed_api_inputs_are_rejected(self):
        self.assertEqual(self.client.post('/api/login', json=[]).status_code, 400)
        self.login()
        for route in ['/api/vault', '/api/permissions', '/api/agents', '/api/tts']:
            self.assertEqual(self.client.post(route, json=[]).status_code, 400)
            self.assertEqual(self.client.post(route, content='{', headers={'Content-Type':'application/json'}).status_code, 400)
        self.assertEqual(self.client.post('/api/agents', json={'tools':5}).status_code, 400)
        self.assertEqual(self.client.post('/api/tts', json={'text':[]}).status_code, 400)

    def test_permissions_enforce_auto_ask_deny_and_reload(self):
        from ali import permissions
        from ali.skills import Skill
        from unittest.mock import Mock
        fn = Mock(return_value='ok')
        tool = Skill('permission_test','Test',{},[],fn,True)
        path=Path(state.name)/'permission-test.json'
        with patch.object(permissions,'PATH',path), patch.dict(brain.registry,{'permission_test':tool}), patch.object(brain,'confirm',return_value=False) as confirm:
            self.assertTrue(brain._run_one('permission_test',{})[1])
            fn.assert_not_called()
            permissions.save({'permission_test':'auto'},brain.registry)
            self.assertEqual(brain._run_one('permission_test',{}),('ok',False))
            self.assertEqual(confirm.call_count,1)
            permissions.save({'permission_test':'deny'},brain.registry)
            self.assertTrue(brain._run_one('permission_test',{})[1])
            self.assertEqual(fn.call_count,1)
            self.assertEqual(permissions.load()['rules']['permission_test'],'deny')
            path.write_text('broken',encoding='utf-8')
            self.assertEqual(permissions.mode('permission_test'),'deny')

    def test_permissions_api_auth_validation_and_catalog(self):
        from ali import permissions
        path=Path(state.name)/'permissions-api.json'
        with patch.object(permissions,'PATH',path):
            self.assertEqual(self.client.get('/api/permissions').status_code,401)
            self.assertEqual(self.client.post('/api/permissions',json={'rules':{}}).status_code,401)
            self.login()
            data=self.client.get('/api/permissions').json()
            self.assertIn('write_project_file',[t['name'] for t in data['tools']])
            self.assertTrue(all('token' not in c and 'secret' not in c for c in data['connections']))
            self.assertEqual(self.client.post('/api/permissions',json={'rules':{'fake_tool':'auto'}}).status_code,400)
            self.assertEqual(self.client.post('/api/permissions',json={'rules':{'web_search':'ask'}}).status_code,400)
            self.assertEqual(self.client.post('/api/permissions',json={'rules':{}},headers={'Origin':'https://untrusted.test'}).status_code,403)
            self.assertEqual(self.client.post('/api/permissions',json={'rules':{'get_datetime':'deny'}}).status_code,200)
            self.assertEqual(permissions.mode('get_datetime'),'deny')

    def test_vault_auth_search_detail_and_persistent_note(self):
        from ali.memory import memory, Memory
        self.assertEqual(self.client.get('/api/vault').status_code, 401)
        self.assertEqual(self.client.post('/api/vault', json={'title':'x','content':'y'}).status_code, 401)
        self.login()
        self.assertEqual(self.client.post('/api/vault', json={'title':'','content':'x'}).status_code, 400)
        response = self.client.post('/api/vault', json={'title':'Vault test', 'content':'Unique vault needle <script>alert(1)</script>'})
        record_id = response.json()['id']
        result = self.client.get('/api/vault', params={'q':'Unique vault needle'}).json()
        self.assertEqual(result['total'],1)
        self.assertEqual(result['items'][0]['id'],record_id)
        self.assertNotIn('content',result['items'][0])
        self.assertIn('<script>',self.client.get('/api/vault/entry',params={'id':record_id}).json()['content'])
        second = Memory(config.db_path)
        try:
            self.assertEqual(second.vault_detail(record_id)['title'],'Vault test')
        finally:
            second.conn.close()
        self.assertEqual(self.client.get('/api/vault/entry',params={'id':'../../.env'}).status_code,404)

    def test_vault_news_history_and_exact_duplicate(self):
        from ali.memory import memory
        one = memory.add_knowledge('[yangiliklar] Test news','Source https://example.com unique-news')
        two = memory.add_knowledge('[yangiliklar] Test news','Source https://example.com unique-news')
        self.assertEqual(one,two)
        memory.log_message('vault-test','user','A past conversation needle')
        self.assertEqual(memory.vault_list('unique-news','news')['total'],1)
        self.assertEqual(memory.vault_list('past conversation needle','chats')['total'],1)
        self.assertEqual(memory.vault_list('unique-news','identity')['total'],0)

    def test_learning_failure_retries_without_claiming_success(self):
        from ali import autopilot
        import threading
        from types import SimpleNamespace
        fake = SimpleNamespace(agents={"tester": SimpleNamespace(learn=[], description="Testing")}, lock=threading.Lock())
        path = Path(state.name) / "learning-test.json"
        with patch.object(autopilot, "STATE", path), patch.object(config, "auto_learn", True), patch.object(config, "watch_projects", False), patch.object(autopilot, "learn", side_effect=RuntimeError("API failed")) as learn:
            autopilot._tick(fake)
            saved = autopilot._load()
            self.assertEqual(saved["learning"]["tester"]["status"], "error")
            self.assertNotIn("tester", saved["learn"])
            self.assertFalse(fake.lock.locked())
            autopilot._tick(fake)
            self.assertEqual(learn.call_count, 1)

    def test_learning_success_tracks_saved_knowledge_and_waits(self):
        from ali import autopilot
        from ali.memory import memory
        import threading
        from types import SimpleNamespace
        fake = SimpleNamespace(agents={"research-test": SimpleNamespace(learn=["Research"], description="Testing")}, lock=threading.Lock())
        path = Path(state.name) / "learning-success.json"
        def learn(*args):
            memory.add_knowledge("[research-test] Verified topic", "Source: https://example.com; checked today")
            return "Saved one topic"
        with patch.object(autopilot, "STATE", path), patch.object(config, "auto_learn", True), patch.object(config, "watch_projects", False), patch.object(autopilot, "learn", side_effect=learn) as run:
            autopilot._tick(fake)
            saved = autopilot.learning_status(fake)
            self.assertEqual(saved["agents"]["research-test"]["status"], "updated")
            self.assertEqual(saved["agents"]["research-test"]["added"], 1)
            autopilot._tick(fake)
            self.assertEqual(run.call_count, 1)

    def test_self_improve_fixes_tools_adds_lessons_and_rolls_back(self):
        from ali import evolution
        from ali.skills import REGISTRY
        agents.save("evo-test", "Evo", "E", "Testing", ["search_knowledge", "missing_tool"], "Base prompt", ["Topic"])
        try:
            report = evolution.health(REGISTRY, "evo-test")
            self.assertIn("missing_tool", report["unknown_tools"])
            self.assertLess(report["score"], 100)
            self.assertEqual(evolution.auto_fix(REGISTRY, "evo-test"), ["Mavjud bo'lmagan vositalar olib tashlandi: missing_tool"])
            self.assertEqual(agents.load()["evo-test"].tools, ["search_knowledge"])
            self.assertIn("Qo'shib bo'lmaydi", evolution.improve(REGISTRY, "evo-test", add_tools=["write_project_file"]))
            result = evolution.improve(REGISTRY, "evo-test", lessons=["Avval manbani tekshir", "Avval manbani tekshir"], learn=["New topic"])
            self.assertIn("+1 saboq", result)
            a = agents.load()["evo-test"]
            self.assertTrue(a.prompt.startswith("Base prompt"))
            self.assertEqual(evolution.split_lessons(a.prompt)[1], ["Avval manbani tekshir"])
            self.assertEqual(a.learn, ["Topic", "New topic"])
            self.assertIn("O'zgarish yo'q", evolution.improve(REGISTRY, "evo-test", lessons=["avval manbani tekshir"]))
            evolution.rollback("evo-test")
            self.assertEqual(agents.load()["evo-test"].prompt, "Base prompt")
            evolution.rollback("evo-test")
            self.assertIn("missing_tool", agents.load()["evo-test"].tools)
        finally:
            agents.delete("evo-test")

    def test_self_improve_cycle_one_agent_per_tick_reports_and_retries(self):
        from ali import autopilot, evolution
        import threading
        from types import SimpleNamespace
        fake = SimpleNamespace(agents={"a1": SimpleNamespace(icon="A", name="One"), "a2": SimpleNamespace(icon="B", name="Two")},
                               lock=threading.Lock(), registry={})
        path = Path(state.name) / "improve-test.json"
        calls = []
        def cycle(brain, key, last):
            calls.append(key)
            if key == "a2" and calls.count("a2") == 1:
                raise RuntimeError("API failed")
            return {"status": "healthy", "score": 100, "issues": [], "fixed": [], "summary": "", "jobs_seen": 0}
        with patch.object(autopilot, "STATE", path), patch.object(config, "auto_learn", False), \
                patch.object(config, "watch_projects", False), patch.object(config, "self_improve", True), \
                patch.object(evolution, "run_cycle_for", side_effect=cycle), patch.object(autopilot.hub, "notify") as notify:
            autopilot._tick(fake)
            autopilot._tick(fake)
            autopilot._tick(fake)
            self.assertEqual(calls, ["ali", "a1", "a2"])
            self.assertEqual(autopilot._load()["improve"]["a2"]["status"], "error")
            self.assertFalse(fake.lock.locked())
            autopilot._tick(fake)  # a2 bir soat kutadi, hisobot hali yo'q
            notify.assert_not_called()
            data = autopilot._load(); data["improve"]["a2"]["retry_at"] = 0; autopilot._save(data)
            autopilot._tick(fake)
            autopilot._tick(fake)
            self.assertEqual(calls, ["ali", "a1", "a2", "a2"])
            notify.assert_called_once()
            self.assertIn("100/100", notify.call_args[0][0])
            self.assertIn("🧠", notify.call_args[0][0])
            autopilot._tick(fake)
            self.assertEqual(len(calls), 4)
            notify.assert_called_once()

    def test_ali_self_improve_lessons_reach_system_prompt_and_roll_back(self):
        from ali import evolution
        from ali.brain import build_system_prompt
        from ali.monitor import monitor
        from ali.skills import REGISTRY
        evolution.ALI_PATH.unlink(missing_ok=True)
        job = monitor.start_job("ali", "Broken request", "test")
        monitor.finish(job, False)
        report = evolution.health(REGISTRY, "ali")
        self.assertEqual(report["key"], "ali")
        self.assertIn("Broken request", report["failed_tasks"])
        self.assertIn("faqat lessons", evolution.improve(REGISTRY, "ali", learn=["x"]))
        self.assertIn("+1 saboq", evolution.improve(REGISTRY, "ali", lessons=["Fayl so'ralsa avval list_projects ishlat"]))
        self.assertIn("Fayl so'ralsa avval list_projects ishlat", build_system_prompt({}))
        self.assertIn("O'zgarish yo'q", evolution.improve(REGISTRY, "ali", lessons=["fayl so'ralsa avval list_projects ishlat"]))
        evolution.rollback("ali")
        self.assertEqual(evolution.ali_lessons(), [])
        self.assertNotIn("Tajriba (o'z xatolaringdan", build_system_prompt({}))

    def test_mtf_signal_blocks_against_higher_timeframe_and_aggregates_4h(self):
        from ali.skills import signals
        hourly = [signals.Bar(i * 3600, 1 + i, 2 + i, i, 1.5 + i) for i in range(8)]
        four = signals._aggregate(hourly, 4 * 3600)
        self.assertEqual([(b.t, b.o, b.h, b.l, b.c) for b in four], [(0, 1, 5, 0, 4.5), (14400, 5, 9, 4, 8.5)])
        up = [signals.Bar(i, 100 + i, 101 + i, 99 + i, 100.5 + i) for i in range(120)]
        down = [signals.Bar(i, 300 - i, 301 - i, 299 - i, 299.5 - i) for i in range(120)]
        fetch = lambda symbol, market, interval, htf=False: down if interval == "1d" or htf else up
        def analyze(symbol, market, interval, bars, htf):
            return signals.Plan(symbol=symbol, market=market, interval=interval, side=1, score=3, setup="EMA cross",
                                entry=1, sl=0.5, tps=(1.75, 2.1, 2.5), zone=(0.9, 1.1))
        with patch.object(signals, "_fetch", side_effect=fetch), patch.object(signals, "analyze", side_effect=analyze):
            plans = signals._mtf_plans("EURUSD", "auto")
            self.assertEqual([p.side for p in plans], [0, 0, 0, 1])
            self.assertIn("Yuqori timeframe qarshi: 1d DOWN", plans[0].reason)
            text = signals.mtf_signal("EURUSD")
            self.assertIn("1d: ▲ LONG", text)
            self.assertIn("Eng kuchli tekshirilgan signal", text)
        def opposed(symbol, market, interval, bars, htf):
            plan = analyze(symbol, market, interval, bars, htf)
            plan.against = 1
            return plan
        with patch.object(signals, "_fetch", side_effect=lambda *a, **k: up), patch.object(signals, "analyze", side_effect=opposed):
            self.assertIn("Hamma tekshiruvdan o'tgan signal yo'q", signals.mtf_signal("EURUSD"))

    def test_mt5_signals_reads_v5_csv_and_hides_unverified(self):
        from ali.skills import signals
        folder = Path(state.name) / "mt5files"
        folder.mkdir(exist_ok=True)
        (folder / "AZT_signals_v5.csv").write_text(
            "time,symbol,timeframe,side,setup,entry,sl,tp1,tp2,tp3,zone_low,zone_high,score,atr,ai,ai_conf,sent,mtf\r\n"
            "2026.10.15 10:00,EURUSD,M15,BUY,EMA cross,1.1,1.09,1.115,1.12,1.13,1.099,1.101,3,0.001,APPROVE,82,1,Yuqori TF: H1 ▲UP\r\n"
            "2026.10.15 11:00,EURUSD,H1,SELL,BOS,1.1,1.11,1.085,1.08,1.07,1.099,1.101,2,0.002,REJECT,40,0,\r\n",
            encoding="utf-8")
        with patch.dict(os.environ, {"MT5_FILES_DIR": str(folder)}):
            sent = signals.mt5_signals()
            self.assertIn("Claude APPROVE 82%", sent)
            self.assertNotIn("SELL", sent)
            self.assertIn("❌ yuborilmagan", signals.mt5_signals(all=True))

    def test_agent_upgrade_keeps_lessons_and_extra_topics(self):
        source, target = Path(state.name) / "upgrade-src", Path(state.name) / "upgrade-dst"
        source.mkdir(exist_ok=True); target.mkdir(exist_ok=True)
        (target / "trader.md").write_text("---\nnom: Old\nbelgi: T\ntavsif: old\nvositalar: a\no'rganish: x; coach topic\n---\n"
                                          "Old prompt\n\n## Tajriba (avtomatik)\n- Keep this lesson\n", encoding="utf-8")
        (source / "trader.md").write_text("---\nnom: New\nbelgi: T\ntavsif: new\nvositalar: a, b\no'rganish: x\nversiya: 2\n---\n"
                                          "New prompt\n", encoding="utf-8")
        self.assertEqual(agents.upgrade_defaults(source, target), ["trader"])
        a = agents._parse("trader", (target / "trader.md").read_text(encoding="utf-8"))
        self.assertTrue(a.prompt.startswith("New prompt"))
        self.assertIn("- Keep this lesson", a.prompt)
        self.assertEqual((a.version, a.tools, a.learn), (2, ["a", "b"], ["x", "coach topic"]))
        self.assertEqual(agents.upgrade_defaults(source, target), [])
        repo = {k: agents._parse(k, (Path(__file__).resolve().parents[1] / "agents" / f"{k}.md").read_text(encoding="utf-8"))
                for k in ("treyder", "dizayner", "office", "informatika")}
        self.assertEqual(repo["dizayner"].version, 2)
        from ali.skills import REGISTRY
        for key in ("dizayner", "office", "informatika"):
            self.assertTrue(set(repo[key].tools) - {"web_search"} <= set(REGISTRY), key)
        self.assertEqual(repo["treyder"].version, 2)
        self.assertTrue(set(repo["treyder"].tools) - {"web_search"} <= set(REGISTRY))

    def test_router_detects_russian_and_office_requests(self):
        from ali import router
        cases = {
            "Сделай презентацию про алгоритмы для 7 класса": ("ru", ["pptx"]),
            "Объясни, что такое цикл for в Python": ("ru", []),
            "Качонки али ботга русча йозсам учитель информатики ишга тушсин": ("uz", []),
            "Менга эхзелда баҳолар жадвалини тузиб бер": ("uz", ["xlsx"]),
            "Python bo'yicha taqdimot tayyorlab ber": ("other", ["pptx"]),
            "dars jadvali qanaqa?": ("other", []),
            "Напиши реферат о сетях в Word": ("ru", ["docx"]),
        }
        for text, (lang, kinds) in cases.items():
            self.assertEqual((router.language(text), router.office_kinds(text)), (lang, kinds), text)
        team = {"informatika": object(), "office": object()}
        both = router.hint("Сделай презентацию про алгоритмы", team)
        self.assertIn("'informatika'", both)
        self.assertIn("'office'", both)
        self.assertIn("'informatika'", router.hint("Объясни рекурсию", team))
        self.assertNotIn("'informatika'", router.hint("Python bo'yicha taqdimot tayyorlab ber", team))
        self.assertEqual(router.hint("Salom, qalaysan?", team), "")
        self.assertEqual(router.hint("Сделай презентацию", {}), "")
        lesson = router.hint("Сделай презентатцию про информации и ворд презентатция наш урок 50 минут",
                             {"informatika": 1, "office": 1, "dizayner": 1})
        self.assertIn("make_lesson_presentation", lesson)
        self.assertIn("with_lesson_plan=true", lesson)
        self.assertNotIn("'office'", lesson)

    def test_brain_prompt_gets_routing_hint_but_memory_keeps_user_text(self):
        from types import SimpleNamespace
        from ali.memory import memory
        fake = SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text="Готово")])
        local = Brain()
        local.agents = {"informatika": SimpleNamespace(prompt="x"), "office": SimpleNamespace(prompt="y")}
        with patch.object(local, "_call", return_value=fake), patch("ali.brain._echo_content", side_effect=lambda c: c), \
                patch.object(local, "_save_session"):
            local.ask("Объясни рекурсию")
        sent = local.messages[-2]["content"]
        self.assertIn("Avtomatik yo'naltirish", sent)
        self.assertTrue(sent.endswith("Объясни рекурсию"))
        self.assertEqual(memory.recent_messages(local.session_id, 2)[0]["content"], "Объясни рекурсию")

    def test_office_documents_are_created_and_reopenable(self):
        from ali import outbox
        from ali.skills import documents, presentation
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        start = outbox.mark()
        result = presentation.make_presentation("Алгоритмы", [
            {"type": "section", "title": "Введение"},
            {"type": "bullets", "title": "Много пунктов", "bullets": [f"Пункт {i}" for i in range(9)], "notes": "Заметка"},
            {"type": "two_column", "title": "Сравнение", "left": {"title": "A", "bullets": ["1"]}, "right": {"title": "B", "bullets": ["2"]}},
            {"type": "table", "title": "Таблица", "rows": [["A", "B"], ["1", "2"]]},
            {"type": "chart", "title": "Диаграмма", "chart": {"type": "pie", "categories": ["x", "y"], "series": [{"name": "s", "values": [1, 2]}]}},
            {"type": "stats", "title": "Цифры", "stats": [{"value": "45", "label": "минут"}]},
            {"type": "quote", "text": "Цитата", "author": "Автор"},
            {"type": "chart", "title": "bad chart", "chart": {}},
        ], subtitle="7 класс", theme="green", ending="Спасибо за внимание!")
        files = outbox.since(start)
        self.assertEqual(len(files), 1)
        deck = Presentation(files[0])
        # title + section + 2 bullet slides (9 > 7) + two_column + table + chart + stats + quote + ending; bad chart skipped
        self.assertEqual(len(deck.slides), 10, result)
        self.assertIn("(2)", " ".join(sh.text_frame.text for sh in deck.slides[3].shapes if sh.has_text_frame))
        self.assertEqual(deck.slides[2].notes_slide.notes_text_frame.text, "Заметка")
        documents.make_word("Реферат", [{"heading": "Введение", "paragraphs": ["Текст"], "bullets": ["a"],
                                         "table": [["A", "B"], ["1", "2"]]}], toc=True)
        doc = Document(outbox.since(start)[1])
        self.assertIn("Введение", [p.text for p in doc.paragraphs])
        self.assertEqual(len(doc.tables), 1)
        documents.make_excel("Оценки", [{"name": "7А", "columns": ["Ученик", "Оценка"], "rows": [["Аня", "5"], ["Боря", "4"],
                                          ["Хак", "=cmd|' /C calc'!A0"]], "totals": True, "chart": {"type": "column", "y_cols": [2]}}])
        ws = load_workbook(outbox.since(start)[2])["7А"]
        self.assertEqual(ws["B2"].value, 5)
        self.assertEqual(ws["B5"].value, "=SUM(B2:B4)")
        self.assertFalse(str(ws["B4"].value).startswith("="))  # DDE formula neutralised
        self.assertEqual(len(ws._charts), 1)
        self.assertIsNone(outbox.resolve("../secret.key"))
        self.assertEqual(outbox.resolve(files[0].name), files[0])

    def test_created_file_reaches_web_and_download_requires_login(self):
        from ali.skills import documents
        def ask(*args, **kwargs):
            documents.make_excel("Test file", [{"columns": ["A"], "rows": [["1"]]}])
            return "Готово"
        with patch("runtime.scheduler.start"), patch("runtime.autopilot.start"), patch.object(brain, "ask", side_effect=ask):
            with TestClient(app) as client:
                client.post("/api/login", json={"password": "test-password"})
                with client.websocket_connect("/ws") as ws:
                    ws.receive_json(); ws.receive_json()
                    ws.send_json({"type": "ask", "id": "file-turn", "text": "Excel", "super": False})
                    for _ in range(15):
                        message = ws.receive_json()
                        if message["type"] == "files":
                            url = message["files"][0]["url"]
                            break
                    else:
                        self.fail("No files event")
                self.assertEqual(client.get(url).status_code, 200)
                client.cookies.clear()
                self.assertEqual(client.get(url).status_code, 401)

    def test_new_default_agent_added_once_and_deleted_stays_deleted(self):
        source, target = Path(state.name) / "defaults-src", Path(state.name) / "defaults-dst"
        source.mkdir(exist_ok=True); target.mkdir(exist_ok=True)
        (source / "dizayner.md").write_text("---\nnom: D\n---\nPrompt\n", encoding="utf-8")
        self.assertEqual(agents.add_new_defaults(source, target), ["dizayner"])
        (target / "dizayner.md").unlink()
        self.assertEqual(agents.add_new_defaults(source, target), [])
        self.assertFalse((target / "dizayner.md").exists())

    def test_lesson_presentation_has_five_timed_stages_homework_and_plan(self):
        from ali import outbox
        from ali.skills import presentation
        from docx import Document
        from pptx import Presentation
        start = outbox.mark()
        stage = lambda n: {"activity": "Действия", "slides": [{"type": "bullets", "title": f"Слайд {i}", "bullets": ["🔍 Пункт"]}
                                                             for i in range(n)]}
        result = presentation.make_lesson_presentation(
            "Поиск информации и кодирование информации", {"organizational": stage(1), "homework_check": stage(1),
             "new_topic": stage(2), "reinforcement": stage(1), "conclusion": stage(1)},
            grade="7", language="ru", goals={"educational": "Цель"}, homework={"tasks": ["§5, вопросы 1-4"]})
        self.assertIn("mazmun kam", result)  # 50 daqiqaga kam slayd haqida ogohlantiradi
        pptx_file, docx_file = outbox.since(start)
        deck = Presentation(pptx_file)
        texts = [" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in deck.slides]
        for name in ("Организационный момент", "Проверка домашнего задания", "Объяснение новой темы",
                     "Закрепление новой темы", "Итог урока", "Домашнее задание", "План урока"):
            self.assertTrue(any(name in t for t in texts), name)
        self.assertTrue(any("Этап 3 · 15 мин" in t for t in texts))
        self.assertEqual(presentation.choose_motif("Поиск информации и кодирование информации"), "binary")
        table = Document(docx_file).tables[0]
        self.assertEqual([table.cell(r, 3).text for r in range(1, 6)], ["5 мин", "10 мин", "15 мин", "15 мин", "5 мин"])
        self.assertEqual(table.cell(1, 2).text, "Действия")

    def test_approved_agent_request_builds_agent_and_notifies(self):
        from types import SimpleNamespace
        from ali import service
        from ali.skills import REGISTRY
        local = Brain()
        local.confirm = lambda name, args: False
        args = {"key": "preapproved-test", "name": "Test", "icon": "T", "description": "d", "tools": [], "prompt": "p"}
        local.preapproved.add("create_agent")
        try:
            out, error = local._run_one("create_agent", args)
            self.assertFalse(error, out)
            out, error = local._run_one("create_agent", args)  # bir martalik: ikkinchisi yana ruxsat so'raydi
            self.assertTrue(error)
            self.assertIn("rad etdi", out)
        finally:
            agents.delete("preapproved-test")

        seen = {}
        def ask(text, *rest):
            seen["preapproved"] = set(brain.preapproved)
            seen["text"] = text
            agents.save("lawyer-test", "Huquqshunos", "⚖️", "Huquqiy savollar", [], "Prompt")
            return "Huquqshunos agenti yaratildi"
        try:
            with patch.object(brain, "ask", side_effect=ask):
                changed, answer = asyncio.run(service.build_agent(7, "Huquqshunos agent kerak"))
            self.assertEqual(changed, ["lawyer-test"])
            self.assertIn("create_agent", seen["preapproved"])
            self.assertIn("Huquqshunos agent kerak", seen["text"])
            self.assertNotIn("create_agent", brain.preapproved)
            bot = TelegramBot()
            bot.owner_id = 1
            bot.app = SimpleNamespace(bot=SimpleNamespace(send_message=AsyncMock()))
            brain.reload_agents()
            with patch.object(service, "build_agent", AsyncMock(return_value=(["lawyer-test"], "Tayyor"))):
                asyncio.run(bot.build_requested_agent({"id": 7, "uid": 42, "text": "Huquqshunos agent kerak"}))
            messages = {call.args[0]: call.args[1] for call in bot.app.bot.send_message.await_args_list}
            self.assertIn("agent tayyor", messages[1])
            self.assertIn("⚖️ Huquqshunos", messages[42])
        finally:
            agents.delete("lawyer-test")
            brain.reload_agents()

    def test_design_floor_plan_home_and_shop_files(self):
        import ezdxf
        from ali import outbox
        from ali.skills import design_floorplan as fp
        start = outbox.mark()
        result = fp.design_floor_plan("Uy 98", "home", area_m2=98)
        files = {p.suffix + ("3d" if "_3d" in p.name else ""): p for p in outbox.since(start)}
        self.assertEqual(set(files), {".png", ".png3d", ".dxf", ".obj3d", ".mtl3d"}, result)
        doc = ezdxf.readfile(files[".dxf"])
        layers = {e.dxf.layer for e in doc.modelspace()}
        self.assertTrue({"A-DEVOR", "A-ESHIK", "A-DERAZA", "A-MEBEL", "A-MATN"} <= layers)
        self.assertGreater(files[".obj3d"].read_text().count("\nv "), 100)
        plan = fp.build_plan("Uy", "home", 98, 0, 0, [], "uz", "grocery", 2.8)
        hall = next(r for r in plan.rooms if r.kind == "hall")
        doors = [o for o in plan.openings if o.kind == "door"]
        # yo'lakli reja: har bir xonaga eshik bor va hech bir xona 1.5 m dan tor emas
        self.assertEqual({id(o.room) for o in doors}, {id(r) for r in plan.rooms if r is not hall})
        self.assertTrue(all(min(r.w, r.h) >= 1.5 for r in plan.rooms if r is not hall))
        self.assertGreaterEqual(min(hall.w, hall.h), 1.2)  # yo'lak me'yori
        self.assertTrue(any(o.kind == "entrance" for o in plan.openings))
        shop = fp.build_plan("Do'kon", "shop", 0, 14, 9, [], "uz", "grocery", 2.8)
        kinds = [f[0] for r in shop.rooms for f in r.furniture]
        self.assertIn("cashier", kinds)
        self.assertGreater(kinds.count("gondola"), 2)
        self.assertTrue(shop.flow and shop.zones)
        for r in shop.rooms:  # jihozlar devordan chiqmaydi va bir-birini kesmaydi
            for i, (_, a, b, c, d) in enumerate(r.furniture):
                self.assertTrue(r.x0 <= a < c <= r.x1 and r.y0 <= b < d <= r.y1)
                for _, a2, b2, c2, d2 in r.furniture[i + 1:]:
                    self.assertFalse(a < c2 and c > a2 and b < d2 and d > b2)

    def test_signage_and_video_outputs(self):
        from ali import outbox
        from ali.skills import design_media
        start = outbox.mark()
        text = design_media.design_signage("MARKET 24", "Oziq-ovqat", width_m=5, height_m=1.2, viewing_distance_m=120)
        names = sorted(p.name for p in outbox.since(start))
        self.assertEqual(len(names), 3)
        self.assertTrue(any(n.endswith("_kun.png") for n in names) and any(n.endswith("_tun.png") for n in names))
        svg = next(p for p in outbox.since(start) if p.suffix == ".svg").read_text()
        self.assertIn('width="5000mm"', svg)
        self.assertIn("harfni kamida", text)  # 120 m uchun 660 mm harf kam
        start = outbox.mark()
        result = design_media.make_promo_video("Test", [{"title": "Salom", "text": "Matn", "duration": 1.5}], "1:1")
        made = outbox.since(start)
        self.assertEqual(sorted(p.suffix for p in made), [".jsx", ".mp4"], result)
        self.assertGreater(made[0].stat().st_size if made[0].suffix == ".mp4" else made[1].stat().st_size, 10000)
        router_team = {"dizayner": 1, "informatika": 1, "office": 1}
        from ali import router
        self.assertIn("'dizayner'", router.hint("80 kvadrat uyimga dizayn qilib ber", router_team))
        self.assertEqual(router.hint("Maximal narx qancha?", router_team), "")

    def test_coach_agent_and_improve_skills_are_registered(self):
        from ali.skills import REGISTRY
        coach = agents._parse("murabbiy", (Path(__file__).resolve().parents[1] / "agents" / "murabbiy.md").read_text(encoding="utf-8"))
        self.assertTrue(set(coach.tools) - {"web_search"} <= set(REGISTRY))
        for name in ("agent_health", "improve_agent", "rollback_agent", "improve_agents_now"):
            self.assertFalse(REGISTRY[name].dangerous)

    def test_restricted_agent_cannot_call_unlisted_tools(self):
        from types import SimpleNamespace
        block = SimpleNamespace(type="tool_use", name="write_project_file", input={}, id="blocked")
        with patch.object(brain, "_run_one") as run:
            results = brain._run_tools([block], lambda *args: None, {"search_knowledge"})
            self.assertTrue(results[0]["is_error"])
            run.assert_not_called()

    def test_agent_respects_web_search_deny_and_sees_user_facts(self):
        from types import SimpleNamespace
        from ali import permissions
        from ali.memory import memory
        captured = {}

        def fake_call(emit, effort, system=None, tools=None, messages=None):
            captured.update(system=system, tools=tools)
            return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text="ok")])

        memory.set_fact("test_sevimli_rang", "ko'k")
        path = Path(state.name) / "permission-web.json"
        try:
            with patch.object(permissions, "PATH", path), patch.object(config, "web_search", True), \
                 patch.object(brain, "_call", side_effect=fake_call):
                permissions.save({"web_search": "deny"}, brain.registry)
                brain._run_agent({"agent": "yangiliklar", "task": "yangilik"}, lambda *a: None)
                self.assertNotIn("web_search", [t["name"] for t in captured["tools"]])
                self.assertIn("test_sevimli_rang", captured["system"])
                permissions.save({"web_search": "auto"}, brain.registry)
                # Oddiy vazifada tashqi qidiruv yo'q - faqat agent yaratish/to'ldirishda
                brain._run_agent({"agent": "yangiliklar", "task": "yangilik"}, lambda *a: None)
                self.assertNotIn("web_search", [t["name"] for t in captured["tools"]])
                brain.research = True
                try:
                    brain._run_agent({"agent": "yangiliklar", "task": "yangilik"}, lambda *a: None)
                finally:
                    brain.research = False
                self.assertIn("web_search", [t["name"] for t in captured["tools"]])
        finally:
            memory.delete_fact("test_sevimli_rang")

    def test_github_push_creates_and_updates_with_sha(self):
        import base64, json
        import httpx
        from ali.skills import github as gh
        seen = []

        def handler(request):
            if request.method == "GET":
                if "missing.py" in request.url.path:
                    return httpx.Response(404, json={})
                return httpx.Response(200, json={"sha": "abc123"})
            seen.append(json.loads(request.content))
            return httpx.Response(201, json={"commit": {"html_url": "https://github.com/u/r/commit/1"}})

        transport = httpx.MockTransport(handler)
        with patch.object(gh, "_client", lambda: httpx.Client(base_url=gh.API, transport=transport)), \
             patch.object(gh, "_repo", lambda r: "u/r"):
            out1 = gh.github_push_file("r", "missing.py", "print(1)", "yangi")
            out2 = gh.github_push_file("r", "ali/x.py", "print(2)", "yangila")
        self.assertIn("commit/1", out1)
        self.assertNotIn("sha", seen[0])
        self.assertEqual(seen[1]["sha"], "abc123")
        self.assertEqual(base64.b64decode(seen[1]["content"]).decode(), "print(2)")
        from ali.skills import REGISTRY
        self.assertTrue(REGISTRY["github_push_file"].dangerous)

    def test_failed_turn_rolls_back_pending_messages(self):
        before = list(brain.messages)
        with patch.object(brain, "_call", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                brain._ask("test savol", None, False, False, None)
        self.assertEqual(brain.messages, before)

    def setUp(self):
        self.client = TestClient(app)

    def login(self):
        result = self.client.post("/api/login", json={"password": "test-password"})
        self.assertEqual(result.status_code, 200)

    def test_site_opens_without_ai_key_and_requires_login(self):
        self.assertEqual(self.client.get("/healthz").status_code, 200)
        self.assertEqual(self.client.get("/").url.path, "/login")
        self.assertEqual(self.client.get("/api/agents").status_code, 401)
        self.login()
        self.assertIn('id="feed"', self.client.get("/").text)
        self.assertEqual(self.client.get("/api/status").status_code, 200)
        self.assertTrue(self.client.get("/api/agents").json()["agents"])

    def test_auth_unicode_is_rejected_without_crashing(self):
        self.assertFalse(auth.check_password("пароль"))
        self.assertFalse(auth.check_token("пароль"))

    def test_session_survives_new_brain(self):
        first = Brain()
        first.messages = [{"role": "user", "content": "salom"},
                          {"role": "assistant", "content": "Assalomu alaykum"}]
        first._save_session()
        second = Brain()
        self.assertEqual(first.session_id, second.session_id)
        self.assertEqual(first.messages, second.messages)
        second.new_session()
        self.assertEqual(Brain().messages, [])

    def test_agent_edits_and_deletion_survive_reload(self):
        agents.save("test-agent", "Test", "T", "Test", [], "Test prompt")
        self.assertIn("test-agent", agents.load())
        agents.delete("test-agent")
        importlib.reload(agents)
        self.assertNotIn("test-agent", agents.load())
        self.assertTrue(str(agents.AGENTS_DIR).startswith(state.name))
        with self.assertRaises(ValueError):
            agents.delete("../secrets")

    def test_websocket_auth_and_history(self):
        from starlette.websockets import WebSocketDisconnect
        with self.assertRaises(WebSocketDisconnect):
            with self.client.websocket_connect("/ws"):
                pass
        self.login()
        with self.client.websocket_connect("/ws") as ws:
            self.assertEqual(ws.receive_json()["type"], "hello")
            self.assertEqual(ws.receive_json()["type"], "history")

    def test_vercel_gateway_does_not_import_runtime(self):
        client = TestClient(gateway)
        with patch.dict(os.environ, {"ALI_BACKEND_URL": ""}):
            self.assertEqual(client.get("/").status_code, 503)
        with patch.dict(os.environ, {"ALI_BACKEND_URL": "https://ali.example.com"}):
            result = client.get("/login?next=chat", follow_redirects=False)
            self.assertEqual(result.status_code, 307)
            self.assertEqual(result.headers["location"], "https://ali.example.com/login?next=chat")
        with patch.dict(os.environ, {"ALI_BACKEND_URL": "https://user:secret@ali.example.com"}):
            self.assertEqual(client.get("/").status_code, 503)

    def test_telegram_panel_and_private_guard(self):
        bot = TelegramBot()
        with patch.object(config, "public_url", "https://ali.example.com"):
            self.assertEqual(bot.site_keyboard().inline_keyboard[0][0].web_app.url,
                             "https://ali.example.com/miniapp")
        from types import SimpleNamespace
        update = SimpleNamespace(effective_chat=SimpleNamespace(type="group"),
                                 effective_message=SimpleNamespace(reply_text=AsyncMock()))
        self.assertFalse(asyncio.run(bot._guard(update)))

    def test_chat_request_reaches_brain_and_returns_response(self):
        with patch("runtime.scheduler.start"), patch("runtime.autopilot.start"), \
                patch.object(brain, "ask", return_value="Sinov javobi") as ask:
            with TestClient(app) as client:
                client.post("/api/login", json={"password": "test-password"})
                with client.websocket_connect("/ws") as ws:
                    ws.receive_json()  # hello
                    ws.receive_json()  # history
                    ws.send_json({"type": "ask", "id": "test-turn", "text": "Salom",
                                  "super": False, "agent": "treyder"})
                    for _ in range(12):
                        message = ws.receive_json()
                        if message["type"] == "done":
                            self.assertEqual(message["text"], "Sinov javobi")
                            self.assertEqual(message["id"], "test-turn")
                            break
                    else:
                        self.fail("Chat did not return a result")
                self.assertEqual(ask.call_args.args[0], "Salom")
                self.assertIs(ask.call_args.args[3], True)
                self.assertIsNone(ask.call_args.args[5])

    def test_only_claude_answers_and_web_search_only_for_agent_build(self):
        from types import SimpleNamespace
        from anthropic.types import TextBlock
        response = SimpleNamespace(stop_reason="end_turn", content=[TextBlock(type="text", text="Toshkent")])
        b = Brain()
        with patch.object(config, "web_search", True):
            b.reload_agents()
        self.assertIn("consult_experts", b.registry)  # GPT + Gemini - Claude o'zi kerak bo'lganda chaqiradi
        with patch.object(b, "_call", return_value=response) as call:
            self.assertEqual(b.ask("O'zbekiston poytaxti qayer?", super_mode=True), "Toshkent")
            self.assertNotIn("web_search", [t["name"] for t in call.call_args.kwargs["tools"]])
            b.ask("Huquqshunos agent yarat", super_mode=True)
            self.assertIn("web_search", [t["name"] for t in call.call_args.kwargs["tools"]])
        self.assertEqual(call.call_count, 2)  # boshqa model (Gemini) chaqirilmaydi


if __name__ == "__main__":
    unittest.main()
